#!/usr/bin/env python3
"""Single-job CUDA transcription with durable progress and original timestamps."""
import argparse
import ctypes
import dataclasses
import datetime as dt
import fcntl
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import signal
import subprocess
import sys
import time

ROOT = Path(__file__).resolve().parent

def stamp(seconds):
    milliseconds = round(seconds * 1000)
    hours, milliseconds = divmod(milliseconds, 3600000)
    minutes, milliseconds = divmod(milliseconds, 60000)
    seconds, milliseconds = divmod(milliseconds, 1000)
    return f'{hours:02}:{minutes:02}:{seconds:02},{milliseconds:03}'

def save_json(path, data):
    temporary = path.with_suffix(path.suffix + '.tmp')
    with temporary.open('w', encoding='utf-8') as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write('\n')
        stream.flush()
        os.fsync(stream.fileno())
    temporary.replace(path)

def plain(value):
    if dataclasses.is_dataclass(value):
        return plain(dataclasses.asdict(value))
    if hasattr(value, '_asdict'):
        return plain(value._asdict())
    if isinstance(value, dict):
        return {k: plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [plain(v) for v in value]
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value

def gpu_state():
    command = ['/usr/lib/wsl/lib/nvidia-smi', '--query-gpu=name,uuid,driver_version,memory.free,memory.used,utilization.gpu', '--format=csv,noheader,nounits']
    for line in subprocess.check_output(command, text=True).splitlines():
        name, uuid, driver, free, used, util = [s.strip() for s in line.split(',')]
        if 'RTX 3090' in name:
            return dict(name=name, uuid=uuid, driver=driver, free_mib=int(free), used_mib=int(used), utilization=int(util))
    raise RuntimeError('RTX 3090 not visible; refusing to select another GPU')

def cuda_free_mib():
    """Check driver allocation headroom as well as WDDM nvidia-smi accounting."""
    cuda = ctypes.CDLL('libcuda.so.1')
    def check(code):
        if code != 0:
            raise RuntimeError(f'CUDA driver admission check failed with code {code}')
    check(cuda.cuInit(0))
    device = ctypes.c_int()
    count = ctypes.c_int()
    check(cuda.cuDeviceGetCount(ctypes.byref(count)))
    for index in range(count.value):
        check(cuda.cuDeviceGet(ctypes.byref(device), index))
        name = ctypes.create_string_buffer(256)
        check(cuda.cuDeviceGetName(name, 256, device))
        if b'RTX 3090' in name.value:
            break
    else:
        raise RuntimeError('CUDA driver cannot find the RTX 3090')
    context = ctypes.c_void_p()
    check(cuda.cuCtxCreate_v2(ctypes.byref(context), 0, device))
    try:
        free, total = ctypes.c_size_t(), ctypes.c_size_t()
        check(cuda.cuMemGetInfo_v2(ctypes.byref(free), ctypes.byref(total)))
        return free.value // (1024 * 1024)
    finally:
        check(cuda.cuCtxDestroy_v2(context))

def validate(segments, duration):
    previous = 0.0
    for segment in segments:
        start, end = segment['start'], segment['end']
        if not (math.isfinite(start) and math.isfinite(end) and 0 <= start <= end <= duration + 0.5 and start >= previous - 0.02):
            raise ValueError(f'Invalid original-audio timestamps: {start}, {end}')
        previous = end

def export(out, segments, metadata):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    validate(segments, metadata['duration_seconds'])
    text = '\n'.join(s['text'] for s in segments) + '\n'
    (out / 'transcript.txt').write_text(text, encoding='utf-8')
    srt = '\n\n'.join(f"{i}\n{stamp(s['start'])} --> {stamp(s['end'])}\n{s['text'].strip()}" for i, s in enumerate(segments, 1)) + '\n'
    (out / 'transcript.srt').write_text(srt, encoding='utf-8')
    save_json(out / 'transcript.json', dict(metadata=metadata, segments=segments))
    title = metadata.get('title') or Path(metadata['input']).stem
    markdown = f'# Transcript\n\nSource: {title}\n\nAutomatic transcript. Speaker names have not been assigned.\n\n'
    markdown += '\n\n'.join(f"**{stamp(s['start']).replace(',', '.')}** {s['text']}" for s in segments)
    (out / 'transcript.md').write_text(markdown + '\n', encoding='utf-8')
    doc = Document()
    # Some python-docx distributions include a bordered Title style template.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.1
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph('Transcript', 'Title')
    doc.add_paragraph(title)
    doc.add_paragraph('Automatic transcript with timestamps from the original recording. Speaker names have not been assigned.')
    paragraphs = []
    for segment in segments:
        if not paragraphs or segment['end'] - paragraphs[-1]['start'] > 60 or len(paragraphs[-1]['text']) > 850:
            paragraphs.append(dict(start=segment['start'], text=segment['text']))
        else:
            paragraphs[-1]['text'] += segment['text']
    for segment in paragraphs:
        p = doc.add_paragraph()
        p.add_run(stamp(segment['start']).replace(',', '.') + '  ').bold = True
        p.add_run(segment['text'])
    for border in doc.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.save(out / 'transcript.docx')

def run(args, model_factory=None):
    source = Path(args.input).expanduser().resolve(strict=True)
    if not source.is_file():
        raise ValueError('Input must be a regular media file')
    out = Path(args.output_dir).expanduser().resolve()
    if out.exists() and any(out.iterdir()):
        raise ValueError('Output directory is occupied; choose a new job/output directory')
    lock = (ROOT / '.gpu.lock').open('a')
    try:
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        raise RuntimeError('Another transcription is running; retry when it finishes') from None
    out.mkdir(parents=True, exist_ok=True)
    started = time.monotonic()
    metadata = dict(status='running', complete=False, pid=os.getpid(), input=str(source), title=args.title or source.stem, started_at=dt.datetime.now(dt.timezone.utc).isoformat(), model=args.model, language_requested=args.language, compute_type=args.compute_type, vad=not args.no_vad, beam_size=5, task='transcribe', initial_prompt=args.initial_prompt, offline=args.offline)
    save_json(out / 'run.json', metadata)
    segments = []
    try:
        probe = json.loads(subprocess.check_output(['ffprobe', '-v', 'error', '-show_format', '-show_streams', '-of', 'json', str(source)], text=True))
        streams = [s for s in probe['streams'] if s['codec_type'] == 'audio']
        if not streams:
            raise ValueError('Input has no audio stream')
        metadata['duration_seconds'] = float(probe['format'].get('duration') or streams[0]['duration'])
        if not math.isfinite(metadata['duration_seconds']) or metadata['duration_seconds'] <= 0:
            raise ValueError('Input duration must be finite and positive')
        with source.open('rb') as stream:
            metadata['input_sha256'] = hashlib.file_digest(stream, 'sha256').hexdigest()
        metadata['versions'] = {p: importlib.metadata.version(p) for p in ['faster-whisper', 'ctranslate2', 'av', 'nvidia-cublas-cu12', 'nvidia-cudnn-cu12', 'python-docx']}
        before = gpu_state()
        os.environ['CUDA_VISIBLE_DEVICES'] = before['uuid']
        free_cuda = cuda_free_mib()
        metadata['gpu_before'] = before
        metadata['cuda_free_mib_before'] = free_cuda
        if min(before['free_mib'], free_cuda) < 10240:
            raise RuntimeError('RTX 3090 has less than 10 GiB free; retry after other work finishes')
        if before['utilization'] > 50:
            raise RuntimeError('RTX 3090 is busy; retry after other work finishes')
        if model_factory is None:
            from faster_whisper import WhisperModel
            model_factory = WhisperModel
        model_path = ROOT / 'models' / args.model
        model_name = str(model_path) if model_path.exists() else args.model
        model = model_factory(model_name, device='cuda', device_index=0, compute_type=args.compute_type, download_root=str(ROOT / 'models'), local_files_only=args.offline)
        receipt = ROOT / 'models' / (args.model + '.receipt.json')
        if receipt.exists():
            metadata['model_receipt'] = json.loads(receipt.read_text())
        result, info = model.transcribe(str(source), language=None if args.language == 'auto' else args.language, task='transcribe', beam_size=5, vad_filter=not args.no_vad, vad_parameters=dict(min_silence_duration_ms=2000) if not args.no_vad else None, word_timestamps=True, initial_prompt=args.initial_prompt)
        metadata['language_detected'] = info.language
        metadata['language_probability'] = info.language_probability
        metadata['duration_after_vad_seconds'] = info.duration_after_vad
        save_json(out / 'run.json', metadata)
        with (out / 'segments.jsonl').open('x', encoding='utf-8') as journal:
            for segment in result:
                item = plain(segment)
                journal.write(json.dumps(item, ensure_ascii=False, allow_nan=False) + '\n')
                journal.flush()
                os.fsync(journal.fileno())
                segments.append(item)
                metadata.update(segment_count=len(segments), last_segment_end=item['end'], elapsed_seconds=round(time.monotonic() - started, 3))
                save_json(out / 'run.json', metadata)
                print(f"{item['end']:.1f}/{metadata['duration_seconds']:.1f}s | {item['text']}", flush=True)
        metadata['segment_count'] = len(segments)
        metadata['elapsed_seconds'] = round(time.monotonic() - started, 3)
        metadata['gpu_after_decode'] = gpu_state()
        metadata['coverage_review'] = dict(first_segment_start=segments[0]['start'] if segments else None, last_segment_end=segments[-1]['end'] if segments else None, audio_quality_reviewed=False, note='Export consistency does not establish speech accuracy or full speech coverage')
        # Exports are final only when run.json.complete becomes true below.
        export(out, segments, metadata)
        metadata.update(status='completed', complete=True, finished_at=dt.datetime.now(dt.timezone.utc).isoformat(), elapsed_seconds=round(time.monotonic() - started, 3))
        save_json(out / 'transcript.json', dict(metadata=metadata, segments=segments))
        save_json(out / 'run.json', metadata)
        print(json.dumps(dict(status='completed', output=str(out), elapsed_seconds=metadata['elapsed_seconds'], segments=len(segments))), flush=True)
    except BaseException as exc:
        metadata.update(status='failed', complete=False, error=f'{type(exc).__name__}: {exc}', elapsed_seconds=round(time.monotonic() - started, 3), segment_count=len(segments))
        save_json(out / 'run.json', metadata)
        raise
    finally:
        lock.close()

def parser():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('input')
    p.add_argument('--output-dir', required=True)
    p.add_argument('--model', default='large-v3', choices=['large-v3'])
    p.add_argument('--language', default='auto', help='auto, en, zh, or another Whisper language code')
    p.add_argument('--compute-type', default='float16', choices=['float16', 'int8_float16'])
    p.add_argument('--no-vad', action='store_true')
    p.add_argument('--initial-prompt')
    p.add_argument('--title', help='Readable source title for document exports')
    p.add_argument('--offline', action='store_true')
    return p

if __name__ == '__main__':
    signal.signal(signal.SIGTERM, lambda *_: sys.exit('Terminated'))
    try:
        run(parser().parse_args())
    except (Exception, KeyboardInterrupt) as exc:
        print(f'Transcription failed: {exc}', file=sys.stderr)
        sys.exit(1)
